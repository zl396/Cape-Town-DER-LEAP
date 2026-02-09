#!/usr/bin/env python3
"""
Generate S3_scenario_assumptions.xlsx from Bass diffusion model output.
This creates a structured Excel workbook for LEAP input data.
"""

import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# Read CSV data
# ============================================================================

csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        's3_bass_diffusion_results.csv')

data = []
with open(csv_path, 'r') as f:
    reader = csv.DictReader(f)
    for row in reader:
        data.append(row)

# ============================================================================
# Style definitions
# ============================================================================

header_font = Font(bold=True, size=11)
title_font = Font(bold=True, size=13)
section_font = Font(bold=True, size=11, color='FFFFFF')
section_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
hi_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
mi_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
li_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

def style_header_row(ws, row, max_col, font=section_font, fill=section_fill):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin_border

def auto_width(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 3, 20)


# ============================================================================
# Create workbook
# ============================================================================

wb = Workbook()

# ============================
# Sheet 1: Adoption Rates
# ============================
ws1 = wb.active
ws1.title = 'Adoption Rates'

ws1.cell(row=1, column=1, value='S3 Progressive Subsidy Scenario - SHS Adoption Rates').font = title_font
ws1.cell(row=2, column=1, value='Bass Diffusion Model Output - Year-by-Year Adoption (%)').font = Font(italic=True)

headers = ['Year',
           'HI Adoption (%)', 'HI HH with SHS', 'HI New Adopters',
           'MI Adoption (%)', 'MI HH with SHS', 'MI New Adopters',
           'LI Adoption (%)', 'LI HH with SHS', 'LI New Adopters',
           'Total HH with SHS']

row = 4
for col_idx, h in enumerate(headers, 1):
    ws1.cell(row=row, column=col_idx, value=h)
style_header_row(ws1, row, len(headers))

prev_hh = {'HI': 0, 'MI': 0, 'LI': 0}
for i, d in enumerate(data):
    r = row + 1 + i
    year = int(d['Year'])
    ws1.cell(row=r, column=1, value=year).border = thin_border

    hi_hh = float(d['HI_HH_with_SHS'])
    mi_hh = float(d['MI_HH_with_SHS'])
    li_hh = float(d['LI_HH_with_SHS'])
    total_hh = float(d['Total_HH_with_SHS'])

    ws1.cell(row=r, column=2, value=round(float(d['HI_Adoption_%']), 4)).border = thin_border
    ws1.cell(row=r, column=3, value=round(hi_hh, 0)).border = thin_border
    ws1.cell(row=r, column=4, value=round(max(0, hi_hh - prev_hh['HI']), 0)).border = thin_border

    ws1.cell(row=r, column=5, value=round(float(d['MI_Adoption_%']), 4)).border = thin_border
    ws1.cell(row=r, column=6, value=round(mi_hh, 0)).border = thin_border
    ws1.cell(row=r, column=7, value=round(max(0, mi_hh - prev_hh['MI']), 0)).border = thin_border

    ws1.cell(row=r, column=8, value=round(float(d['LI_Adoption_%']), 4)).border = thin_border
    ws1.cell(row=r, column=9, value=round(li_hh, 0)).border = thin_border
    ws1.cell(row=r, column=10, value=round(max(0, li_hh - prev_hh['LI']), 0)).border = thin_border

    ws1.cell(row=r, column=11, value=round(total_hh, 0)).border = thin_border

    # Apply income group colors
    for c in [2, 3, 4]:
        ws1.cell(row=r, column=c).fill = hi_fill
    for c in [5, 6, 7]:
        ws1.cell(row=r, column=c).fill = mi_fill
    for c in [8, 9, 10]:
        ws1.cell(row=r, column=c).fill = li_fill

    prev_hh = {'HI': hi_hh, 'MI': mi_hh, 'LI': li_hh}

auto_width(ws1)

# ============================
# Sheet 2: SHS Capacity
# ============================
ws2 = wb.create_sheet('SHS Capacity')

ws2.cell(row=1, column=1, value='S3 Progressive Subsidy - SHS Capacity (MW)').font = title_font
ws2.cell(row=2, column=1, value='Residential SHS capacity by income group').font = Font(italic=True)

headers2 = ['Year',
            'HI Capacity (MW)', 'MI Capacity (MW)', 'LI Capacity (MW)',
            'Total Residential (MW)',
            'BAU SSEG (MW) est.', 'S3 Total SSEG (MW) est.']

row = 4
for col_idx, h in enumerate(headers2, 1):
    ws2.cell(row=row, column=col_idx, value=h)
style_header_row(ws2, row, len(headers2))

# BAU SSEG at 8% growth from 121 MW (2023)
bau_sseg_2023 = 121.0
bau_growth = 0.08

for i, d in enumerate(data):
    r = row + 1 + i
    year = int(d['Year'])

    hi_cap = float(d['HI_Capacity_MW'])
    mi_cap = float(d['MI_Capacity_MW'])
    li_cap = float(d['LI_Capacity_MW'])
    total_res = float(d['Total_Residential_Capacity_MW'])

    # BAU SSEG estimate (8% growth, includes commercial)
    years_from_2023 = year - 2023
    bau_sseg = bau_sseg_2023 * (1 + bau_growth) ** years_from_2023 if years_from_2023 >= 0 else bau_sseg_2023

    # S3 total: residential from Bass model + commercial at BAU rate
    # Assume commercial is ~30% of BAU SSEG
    commercial_sseg = bau_sseg * 0.30
    s3_total = total_res + commercial_sseg

    ws2.cell(row=r, column=1, value=year).border = thin_border
    ws2.cell(row=r, column=2, value=round(hi_cap, 2)).border = thin_border
    ws2.cell(row=r, column=3, value=round(mi_cap, 2)).border = thin_border
    ws2.cell(row=r, column=4, value=round(li_cap, 2)).border = thin_border
    ws2.cell(row=r, column=5, value=round(total_res, 2)).border = thin_border
    ws2.cell(row=r, column=6, value=round(bau_sseg, 2)).border = thin_border
    ws2.cell(row=r, column=7, value=round(s3_total, 2)).border = thin_border

    ws2.cell(row=r, column=2).fill = hi_fill
    ws2.cell(row=r, column=3).fill = mi_fill
    ws2.cell(row=r, column=4).fill = li_fill

auto_width(ws2)

# ============================
# Sheet 3: Energy Parameters
# ============================
ws3 = wb.create_sheet('Energy Parameters')

ws3.cell(row=1, column=1, value='S3 Energy Parameters for LEAP Input').font = title_font
ws3.cell(row=2, column=1, value='Grid vs SHS energy consumption split').font = Font(italic=True)

# Energy intensity data (from report Table 8, 2018 values)
ws3.cell(row=4, column=1, value='Baseline Energy Intensity (kWh/HH/year)').font = header_font
params_data = [
    ['', 'Grid Only', 'Grid+SHS (Grid portion)', 'Grid+SHS (SHS portion)', 'SHS Self-Sufficiency %'],
    ['High-Income', 6533, 5492, 1041, '16%'],
    ['Middle-Income', 6347, 5020, 1326, '21%'],
    ['Low-Income', 3887, 3303, 584, '15%'],
]
for i, prow in enumerate(params_data):
    for j, val in enumerate(prow):
        cell = ws3.cell(row=5+i, column=1+j, value=val)
        cell.border = thin_border
        if i == 0:
            cell.font = header_font
            cell.fill = section_fill
            cell.font = section_font

# Notes
ws3.cell(row=11, column=1, value='Notes:').font = header_font
ws3.cell(row=12, column=1, value='- Grid+SHS (Grid portion): consumption from grid for households with SHS')
ws3.cell(row=13, column=1, value='- Grid+SHS (SHS portion): consumption from SHS for households with SHS')
ws3.cell(row=14, column=1, value='- SHS Self-Sufficiency: % of household consumption met by SHS')
ws3.cell(row=15, column=1, value='- Low-income SHS self-sufficiency estimated at 15% (RES4Africa baseline)')
ws3.cell(row=16, column=1, value='- These values should be verified against actual LEAP model parameters')
ws3.cell(row=18, column=1, value='Average SHS System Size (kW)').font = header_font
ws3.cell(row=19, column=1, value='High-Income: 8 kW')
ws3.cell(row=20, column=1, value='Middle-Income: 5 kW')
ws3.cell(row=21, column=1, value='Low-Income: 3 kW')

ws3.cell(row=23, column=1, value='SSEG Capacity Factor: 19.6% (Cape Town SOEC 2021)').font = header_font
ws3.cell(row=24, column=1, value='Household Growth Rate: 1.3% CAGR (IRP projection)')

auto_width(ws3)

# ============================
# Sheet 4: Bass Model Parameters
# ============================
ws4 = wb.create_sheet('Bass Model Parameters')

ws4.cell(row=1, column=1, value='Bass Diffusion Model Parameters & Sensitivity').font = title_font
ws4.cell(row=2, column=1, value='Calibration mode: hybrid (p from data, q from literature)').font = Font(italic=True)

# Parameter table
ws4.cell(row=4, column=1, value='Calibrated Parameters').font = header_font
param_headers = ['Income Group', 'p (innovation)', 'q (imitation)', 'm (saturation)', 'Y0 (2023)', 'R²', 'Policy Mechanism']
for j, h in enumerate(param_headers):
    cell = ws4.cell(row=5, column=1+j, value=h)
    cell.font = section_font
    cell.fill = section_fill
    cell.border = thin_border

param_rows = [
    ['High-Income', 0.027055, 0.38, '30%', '8.00%', 0.508, 'No additional incentive (BAU)'],
    ['Middle-Income', 0.040583, 0.30, '25%', '0.18%', 'N/A (derived)', 'Moderate progressive subsidy'],
    ['Low-Income', 0.067638, 0.19, '15%', '0%', 'N/A (derived)', 'Largest progressive subsidy'],
]
for i, prow in enumerate(param_rows):
    fills = [hi_fill, mi_fill, li_fill]
    for j, val in enumerate(prow):
        cell = ws4.cell(row=6+i, column=1+j, value=val)
        cell.border = thin_border
        cell.fill = fills[i]

# Policy multipliers
ws4.cell(row=10, column=1, value='Policy Multipliers (relative to calibrated HI)').font = header_font
mult_data = [
    ['', 'p multiplier', 'q source'],
    ['High-Income', '1.0x (baseline)', 'Literature: 0.38'],
    ['Middle-Income', '1.5x (moderate subsidy)', 'Literature: 0.30'],
    ['Low-Income', '2.5x (largest subsidy)', 'Literature: 0.19'],
]
for i, mrow in enumerate(mult_data):
    for j, val in enumerate(mrow):
        cell = ws4.cell(row=11+i, column=1+j, value=val)
        cell.border = thin_border
        if i == 0:
            cell.font = header_font

# Sensitivity results
ws4.cell(row=17, column=1, value='Sensitivity Analysis (±30% variation in p and q)').font = header_font
sens_headers = ['Income Group', '2050 Low (-30%)', '2050 Baseline', '2050 High (+30%)']
for j, h in enumerate(sens_headers):
    cell = ws4.cell(row=18, column=1+j, value=h)
    cell.font = section_font
    cell.fill = section_fill
    cell.border = thin_border

# These values are from the script output
sens_rows = [
    ['High-Income', '30.0%', '30.0%', '30.0%'],
    ['Middle-Income', '24.7%', '25.0%', '25.0%'],
    ['Low-Income', '14.6%', '15.0%', '15.0%'],
]
for i, srow in enumerate(sens_rows):
    for j, val in enumerate(srow):
        cell = ws4.cell(row=19+i, column=1+j, value=val)
        cell.border = thin_border

# Literature references
ws4.cell(row=23, column=1, value='Literature References for q Values').font = header_font
ws4.cell(row=24, column=1, value='- Batista da Silva et al.: Market diffusion of household PV systems (Bass model)')
ws4.cell(row=25, column=1, value='- Mejia 2024: Spatiotemporal Estimation of PV Adoption')
ws4.cell(row=26, column=1, value='- Rogers (2003): Diffusion of Innovations, 5th ed.')
ws4.cell(row=27, column=1, value='- Schilling et al. 2009: Technology S-curves in renewable energy')

auto_width(ws4)

# ============================
# Save
# ============================
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           'S3_scenario_assumptions.xlsx')
wb.save(output_path)
print(f"Excel file saved to: {output_path}")
